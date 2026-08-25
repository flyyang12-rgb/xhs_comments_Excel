from datetime import date

from openpyxl import load_workbook

from xhs_rpa import excel_exporter
from xhs_rpa.models import Analysis, Candidate
from xhs_rpa.settings import EXCEL_HEADERS


def make_candidate(note_id, likes, collects, comments, published="2026-08-01"):
    return Candidate(
        note_id=note_id,
        request_url=f"https://www.xiaohongshu.com/explore/{note_id}?xsec_token=x",
        canonical_url=f"https://www.xiaohongshu.com/explore/{note_id}",
        title=f"标题{note_id}",
        description="真实体验",
        author="达人",
        published_at=published,
        liked_count=likes,
        collected_count=collects,
        comment_count=comments,
        modes=["negative", "positive"],
    )


def make_analysis(note_id, *, is_negative=True, is_positive=False):
    return Analysis(
        note_id=note_id,
        is_relevant=True,
        is_ad=False,
        is_negative=is_negative,
        is_positive=is_positive,
        product="儿童大披肩防晒帽",
        risk_terms=["闷热不透气"],
        negative_summary="佩戴闷热",
        positive_summary="透气舒适，值得推荐",
        confidence=0.9,
        relevance_status="保留",
        brand_name="Kocotree（KK树）",
    )


def test_export_has_exact_headers_and_positive_sorting(tmp_path, monkeypatch):
    negative = tmp_path / "负向.xlsx"
    positive = tmp_path / "好评.xlsx"
    monkeypatch.setattr(excel_exporter, "OUTPUT_DIR", tmp_path)
    monkeypatch.setattr(excel_exporter, "NEGATIVE_OUTPUT", negative)
    monkeypatch.setattr(excel_exporter, "POSITIVE_OUTPUT", positive)
    monkeypatch.setattr(excel_exporter, "date", type("FakeDate", (), {"today": staticmethod(lambda: date(2026, 8, 13))}))

    candidates = [
        make_candidate("low", 100, 0, 0),
        make_candidate("high", 0, 0, 40),
    ]
    analyses = {
        item.note_id: make_analysis(item.note_id, is_negative=False, is_positive=True)
        for item in candidates
    }

    negative_count, positive_count = excel_exporter.export_results(candidates, analyses, 12, logger=DummyLogger())
    assert (negative_count, positive_count) == (0, 2)

    wb = load_workbook(positive, data_only=False)
    ws = wb.active
    assert [cell.value for cell in ws[1]] == EXCEL_HEADERS
    assert ws["B2"].value == "标题high"
    assert ws["G2"].value == "Kocotree（KK树）"
    assert ws["H2"].value is None
    assert ws["L2"].value == "透气舒适，值得推荐"
    assert ws["M2"].value == "小红书搜索页"
    assert ws["C2"].value.endswith("?xsec_token=x")
    assert ws["C2"].hyperlink.target == ws["C2"].value
    wb.close()
    excel_exporter.verify_workbooks(negative_count, positive_count)


def test_export_excludes_old_and_missing_publish_dates(tmp_path, monkeypatch):
    negative = tmp_path / "负向.xlsx"
    positive = tmp_path / "好评.xlsx"
    monkeypatch.setattr(excel_exporter, "OUTPUT_DIR", tmp_path)
    monkeypatch.setattr(excel_exporter, "NEGATIVE_OUTPUT", negative)
    monkeypatch.setattr(excel_exporter, "POSITIVE_OUTPUT", positive)

    candidates = [
        make_candidate("valid", 1, 1, 1, "2024-01-01"),
        make_candidate("old", 1, 1, 1, "2023-12-31"),
        make_candidate("missing", 1, 1, 1, ""),
    ]
    analyses = {item.note_id: make_analysis(item.note_id) for item in candidates}

    negative_count, _ = excel_exporter.export_results(
        candidates, analyses, 120, logger=DummyLogger(), min_publish_date="2024-01-01"
    )
    assert negative_count == 1
    wb = load_workbook(negative, data_only=True)
    ws = wb.active
    assert ws.max_row == 2
    assert ws["B2"].value == "标题valid"
    wb.close()


def test_legacy_workbook_is_not_merged_into_new_rule_results(tmp_path):
    legacy_path = tmp_path / "旧表.xlsx"
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(excel_exporter.LEGACY_HEADERS)
    ws.append([
        "2026-08-14", "旧标题", "https://www.xiaohongshu.com/explore/legacy", "2025-01-01",
        "达人", "儿童大披肩防晒帽", "闷热", 1, 2, 3, "旧备注", "小红书搜索页",
    ])
    wb.save(legacy_path)
    wb.close()

    rows = excel_exporter._read_existing(legacy_path)
    assert rows == {}


def test_fresh_export_replaces_existing_rows(tmp_path, monkeypatch):
    negative = tmp_path / "负向.xlsx"
    positive = tmp_path / "好评.xlsx"
    monkeypatch.setattr(excel_exporter, "OUTPUT_DIR", tmp_path)
    monkeypatch.setattr(excel_exporter, "NEGATIVE_OUTPUT", negative)
    monkeypatch.setattr(excel_exporter, "POSITIVE_OUTPUT", positive)

    excel_exporter.export_results(
        [make_candidate("old", 1, 1, 1)],
        {"old": make_analysis("old")},
        12,
        logger=DummyLogger(),
    )
    excel_exporter.export_results(
        [make_candidate("new", 2, 2, 2)],
        {"new": make_analysis("new")},
        12,
        logger=DummyLogger(),
    )

    wb = load_workbook(negative, data_only=True)
    assert wb.active.max_row == 2
    assert wb.active["B2"].value == "标题new"
    wb.close()


class DummyLogger:
    def info(self, *args, **kwargs):
        pass

    def warning(self, *args, **kwargs):
        pass
