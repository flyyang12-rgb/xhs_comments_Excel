from __future__ import annotations

import random
import sys
import time
from pathlib import Path
from urllib.parse import quote

from openpyxl import load_workbook

from xhs_rpa.collector import SpiderXHSCollector
from xhs_rpa.normalization import canonical_note_url
from xhs_rpa.runtime import configure_logging, ensure_runtime_dirs, load_json, save_json
from xhs_rpa.settings import CACHE_DIR, NEGATIVE_OUTPUT, POSITIVE_OUTPUT, VENDOR_ROOT


LINK_REFRESH_CACHE = CACHE_DIR / "link_refresh_state.json"


def _note_id(url: str) -> str:
    canonical = canonical_note_url(url)
    return canonical.rstrip("/").split("/")[-1] if canonical else ""


def _load_targets() -> dict[str, list[tuple[Path, int, str]]]:
    targets: dict[str, list[tuple[Path, int, str]]] = {}
    for path in (NEGATIVE_OUTPUT, POSITIVE_OUTPUT):
        workbook = load_workbook(path, data_only=False, read_only=False)
        sheet = workbook.active
        for row_index in range(2, sheet.max_row + 1):
            url = str(sheet.cell(row_index, 3).value or "")
            if "xsec_token=" in url:
                continue
            note_id = _note_id(url)
            title = str(sheet.cell(row_index, 2).value or "").strip()
            if note_id and title:
                targets.setdefault(note_id, []).append((path, row_index, title))
        workbook.close()
    return targets


def _search_access_url(api, note_id: str, title: str) -> str:
    vendor = str(VENDOR_ROOT)
    if vendor not in sys.path:
        sys.path.insert(0, vendor)
    from xhs_utils.xhs_pc.params import generate_search_id

    search_id = generate_search_id()
    query = title[:80]
    for page in (1, 2):
        success, _, response = api.search_note(
            query,
            page=page,
            sort_type_choice=0,
            note_type=0,
            note_time=0,
            search_id=search_id,
        )
        if not success:
            continue
        data = (response or {}).get("data") or {}
        for item in data.get("items") or []:
            if str(item.get("id") or "") != note_id:
                continue
            token = str(item.get("xsec_token") or "").strip()
            if token:
                return (
                    f"https://www.xiaohongshu.com/explore/{note_id}"
                    f"?xsec_token={token}&xsec_source=pc_search"
                )
        if not data.get("has_more"):
            break
        time.sleep(random.uniform(1.0, 1.8))
    return ""


def _apply_links(
    targets: dict[str, list[tuple[Path, int, str]]],
    access_urls: dict[str, str],
) -> tuple[int, int]:
    found = 0
    fallback = 0
    for path in (NEGATIVE_OUTPUT, POSITIVE_OUTPUT):
        workbook = load_workbook(path, data_only=False, read_only=False)
        sheet = workbook.active
        for note_id, locations in targets.items():
            for target_path, row_index, title in locations:
                if target_path != path:
                    continue
                cell = sheet.cell(row_index, 3)
                access_url = access_urls.get(note_id, "")
                if access_url:
                    cell.value = access_url
                    cell.hyperlink = access_url
                    found += 1
                else:
                    search_url = (
                        "https://www.xiaohongshu.com/search_result?keyword="
                        f"{quote(title)}&source=web_search_result_notes"
                    )
                    cell.hyperlink = search_url
                    fallback += 1
                cell.style = "Hyperlink"
        workbook.save(path)
        workbook.close()
    return found, fallback


def main() -> int:
    ensure_runtime_dirs()
    logger = configure_logging()
    targets = _load_targets()
    if not targets:
        print("所有笔记已包含电脑端访问参数。")
        return 0

    print(f"待定向补链的旧笔记：{len(targets)} 篇")
    collector = SpiderXHSCollector({"throttle": {}}, logger)
    collector.login()

    access_urls: dict[str, str] = load_json(LINK_REFRESH_CACHE, {})
    for index, (note_id, locations) in enumerate(targets.items(), 1):
        if note_id in access_urls:
            continue
        title = locations[0][2]
        print(f"[{index}/{len(targets)}] {title[:36]}")
        try:
            access_url = _search_access_url(collector.api, note_id, title)
        except Exception as error:  # noqa: BLE001 - 单篇失败不中断整批
            logger.warning("定向补链失败 %s：%s", note_id, error)
            access_url = ""
        if access_url:
            access_urls[note_id] = access_url
            save_json(LINK_REFRESH_CACHE, access_urls)
        if index < len(targets):
            time.sleep(random.uniform(1.8, 3.0))

    found, fallback = _apply_links(targets, access_urls)
    if LINK_REFRESH_CACHE.exists():
        LINK_REFRESH_CACHE.unlink()
    print(f"补链完成：直达链接 {found} 个单元格，搜索兜底 {fallback} 个单元格。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
